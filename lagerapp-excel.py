"""
LAGERHANTERING - EXCEL-VERSION
Stöder Excel-filer för import
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import sqlite3
import pandas as pd
import os
from datetime import datetime
import openpyxl

class LagerSystemExcel:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Lagerhantering - Excel Version")
        self.root.geometry("1300x750")
        
        # Databas
        self.conn = sqlite3.connect("lager_excel.db")
        self.cursor = self.conn.cursor()
        self.skapa_tabeller()
        
        # GUI
        self.skapa_gui()
        self.root.mainloop()
    
    def skapa_tabeller(self):
        # Masterdata - produktinformation
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS masterdata (
                sku TEXT PRIMARY KEY,
                artikelnummer TEXT UNIQUE,
                benamning TEXT NOT NULL,
                kategori TEXT,
                underkategori TEXT,
                enhet TEXT DEFAULT 'st',
                inkopspris REAL,
                forsaljningspris REAL,
                mervardesskatt REAL DEFAULT 0.25,
                leverantor TEXT,
                min_niva INTEGER DEFAULT 5,
                max_niva INTEGER DEFAULT 100,
                lead_time INTEGER DEFAULT 7,
                created_date DATE DEFAULT CURRENT_DATE
            )
        ''')
        
        # Stock - aktuellt lager
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS stock (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sku TEXT,
                lagerplats TEXT,
                hylla TEXT,
                rad TEXT,
                fack TEXT,
                batch TEXT,
                serienummer TEXT,
                utgangsdatum DATE,
                antal INTEGER DEFAULT 0,
                reserverat INTEGER DEFAULT 0,
                tillgangligt INTEGER GENERATED ALWAYS AS (antal - reserverat) VIRTUAL,
                senast_raknad DATE,
                status TEXT DEFAULT 'Aktiv',
                kvalitetskod TEXT DEFAULT 'OK',
                anmarkning TEXT,
                FOREIGN KEY (sku) REFERENCES masterdata(sku)
            )
        ''')
        
        # Inbound - inkommande ordrar
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS inbound (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ordernummer TEXT,
                leverantorsorder TEXT,
                sku TEXT,
                leverantor TEXT,
                antal_orderat INTEGER,
                antal_mottaget INTEGER DEFAULT 0,
                enhetspris REAL,
                valuta TEXT DEFAULT 'SEK',
                forvantad_datum DATE,
                mottagen_datum DATE,
                fakturanummer TEXT,
                status TEXT DEFAULT 'På väg',
                anmarkning TEXT,
                FOREIGN KEY (sku) REFERENCES masterdata(sku)
            )
        ''')
        
        # Outbound - utgående ordrar
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS outbound (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                ordernummer TEXT,
                kundorder TEXT,
                sku TEXT,
                kund TEXT,
                kundnummer TEXT,
                antal_bestallt INTEGER,
                antal_levererat INTEGER DEFAULT 0,
                enhetspris REAL,
                valuta TEXT DEFAULT 'SEK',
                ordrad_status TEXT DEFAULT 'Öppen',
                planerat_datum DATE,
                skickat_datum DATE,
                fraktbolag TEXT,
                fraktnummer TEXT,
                anmarkning TEXT,
                FOREIGN KEY (sku) REFERENCES masterdata(sku)
            )
        ''')
        
        # Vyn för lagerstatus med beräknade fält
        self.cursor.execute('''
            CREATE VIEW IF NOT EXISTS v_lagerstatus AS
            SELECT 
                m.sku,
                m.artikelnummer,
                m.benamning,
                m.kategori,
                m.underkategori,
                m.enhet,
                m.forsaljningspris,
                COALESCE(SUM(s.antal), 0) as totalt_antal,
                COALESCE(SUM(s.reserverat), 0) as totalt_reserverat,
                COALESCE(SUM(s.tillgangligt), 0) as tillgangligt,
                m.min_niva,
                m.max_niva,
                (COALESCE(SUM(s.tillgangligt), 0) * m.forsaljningspris) as lager_varde,
                ROUND((COALESCE(SUM(s.tillgangligt), 0) * 100.0 / NULLIF(m.max_niva, 0)), 1) as fyllnadsgrad,
                CASE 
                    WHEN COALESCE(SUM(s.tillgangligt), 0) = 0 THEN '❌ SLUT'
                    WHEN COALESCE(SUM(s.tillgangligt), 0) < m.min_niva THEN '⚠️ LÅGT'
                    WHEN COALESCE(SUM(s.tillgangligt), 0) > m.max_niva * 1.5 THEN '📦 OVERSTOCK'
                    WHEN julianday('now') - julianday(MAX(s.senast_raknad)) > 90 THEN '💀 DEADSTOCK'
                    ELSE '✅ OK'
                END as status,
                MAX(s.senast_raknad) as senast_raknad,
                julianday('now') - julianday(MAX(s.senast_raknad)) as dagar_sen_sista_rorelse
            FROM masterdata m
            LEFT JOIN stock s ON m.sku = s.sku
            GROUP BY m.sku, m.artikelnummer, m.benamning
        ''')
        
        self.conn.commit()
    
    def skapa_gui(self):
        # Notebook (flikar)
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        # Statusrad (skapas tidigt så andra metoder kan uppdatera den)
        self.status_var = tk.StringVar()
        self.status_var.set("Klar")
        status_bar = ttk.Label(self.root, textvariable=self.status_var, 
                              relief=tk.SUNKEN, anchor=tk.W)
        status_bar.pack(side=tk.BOTTOM, fill=tk.X)

        # Skapa flikar
        self.tab_lager = ttk.Frame(self.notebook)
        self.tab_import = ttk.Frame(self.notebook)
        self.tab_rapporter = ttk.Frame(self.notebook)
        self.tab_admin = ttk.Frame(self.notebook)
        
        self.notebook.add(self.tab_lager, text="📦 Lageröversikt")
        self.notebook.add(self.tab_import, text="📁 Excel Import")
        self.notebook.add(self.tab_rapporter, text="📊 Rapporter")
        self.notebook.add(self.tab_admin, text="⚙️ Administration")
        
        # Fyll flikarna
        self.skapa_lager_tab()
        self.skapa_import_tab()
        self.skapa_rapporter_tab()
        self.skapa_admin_tab()
    
    def skapa_lager_tab(self):
        # Sökruta och filter
        filter_frame = ttk.Frame(self.tab_lager)
        filter_frame.pack(fill=tk.X, padx=10, pady=10)
        
        ttk.Label(filter_frame, text="Sök:").pack(side=tk.LEFT, padx=5)
        self.search_var = tk.StringVar()
        search_entry = ttk.Entry(filter_frame, textvariable=self.search_var, width=30)
        search_entry.pack(side=tk.LEFT, padx=5)
        
        ttk.Button(filter_frame, text="Sök", command=self.sok_lager).pack(side=tk.LEFT, padx=5)
        ttk.Button(filter_frame, text="Rensa", command=lambda: [self.search_var.set(""), self.sok_lager()]).pack(side=tk.LEFT, padx=5)
        
        # Filter för status
        ttk.Label(filter_frame, text="Filter:").pack(side=tk.LEFT, padx=(20,5))
        self.filter_var = tk.StringVar(value="ALLA")
        
        for text, value in [("Alla", "ALLA"), ("Lågt", "LÅGT"), ("Overstock", "OVERSTOCK"), ("Deadstock", "DEADSTOCK"), ("Slut", "SLUT")]:
            ttk.Radiobutton(filter_frame, text=text, variable=self.filter_var, 
                           value=value, command=self.sok_lager).pack(side=tk.LEFT, padx=2)
        
        # Trädvy för lager
        columns = ("SKU", "Artikelnummer", "Benämning", "Kategori", "Totalt", 
                  "Reserverat", "Tillgängligt", "Min", "Max", "Värde", "Status")
        
        self.tree = ttk.Treeview(self.tab_lager, columns=columns, show="headings", height=25)
        
        # Definiera kolumnbredder
        col_widths = [120, 120, 250, 100, 80, 80, 80, 60, 60, 100, 100]
        for i, col in enumerate(columns):
            self.tree.heading(col, text=col)
            self.tree.column(col, width=col_widths[i], minwidth=50)
        
        # Scrollbars
        y_scrollbar = ttk.Scrollbar(self.tab_lager, orient=tk.VERTICAL, command=self.tree.yview)
        x_scrollbar = ttk.Scrollbar(self.tab_lager, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=y_scrollbar.set, xscrollcommand=x_scrollbar.set)
        
        # Placering
        self.tree.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=10, pady=(0,10))
        y_scrollbar.pack(side=tk.RIGHT, fill=tk.Y, pady=(0,30))
        x_scrollbar.pack(side=tk.BOTTOM, fill=tk.X, padx=10)
        
        # Ladda data
        self.ladda_lager_data()
        
        # Dubbelklick för detaljer
        self.tree.bind("<Double-1>", self.visa_produktdetaljer)
    
    def ladda_lager_data(self, sok_term="", filter_status="ALLA"):
        """Ladda lagerdata från databasen"""
        # Rensa träd
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        # Bygg SQL-fråga baserat på sökterm och filter
        query = '''
            SELECT sku, artikelnummer, benamning, kategori, totalt_antal, 
                   totalt_reserverat, tillgangligt, min_niva, max_niva, 
                   lager_varde, status
            FROM v_lagerstatus
            WHERE 1=1
        '''
        params = []
        
        if sok_term:
            query += " AND (sku LIKE ? OR artikelnummer LIKE ? OR benamning LIKE ? OR kategori LIKE ?)"
            params.extend([f"%{sok_term}%", f"%{sok_term}%", f"%{sok_term}%", f"%{sok_term}%"])
        
        if filter_status != "ALLA":
            if filter_status == "LÅGT":
                query += " AND status LIKE '%LÅGT%'"
            elif filter_status == "OVERSTOCK":
                query += " AND status LIKE '%OVERSTOCK%'"
            elif filter_status == "DEADSTOCK":
                query += " AND status LIKE '%DEADSTOCK%'"
            elif filter_status == "SLUT":
                query += " AND status LIKE '%SLUT%'"
        
        query += " ORDER BY sku"
        
        self.cursor.execute(query, params)
        
        for row in self.cursor.fetchall():
            # Formatera värde
            formaterad_rad = list(row)
            if formaterad_rad[9]:  # lager_varde
                formaterad_rad[9] = f"{formaterad_rad[9]:,.0f} SEK"
            
            # Bestäm färgkod
            tags = ()
            if "LÅGT" in str(row[10]):
                tags = ('lagt',)
            elif "OVERSTOCK" in str(row[10]):
                tags = ('overstock',)
            elif "DEADSTOCK" in str(row[10]):
                tags = ('deadstock',)
            elif "SLUT" in str(row[10]):
                tags = ('slut',)
            
            self.tree.insert("", "end", values=formaterad_rad, tags=tags)
        
        # Konfigurera färger
        self.tree.tag_configure('lagt', background='#ffcccc')
        self.tree.tag_configure('overstock', background='#fff3cd')
        self.tree.tag_configure('deadstock', background='#f8d7da')
        self.tree.tag_configure('slut', background='#d3d3d3')
        
        # Uppdatera status
        antal_produkter = len(self.tree.get_children())
        self.status_var.set(f"Visar {antal_produkter} produkter")
    
    def sok_lager(self):
        """Utför sökning baserat på sökruta och filter"""
        sok_term = self.search_var.get()
        filter_status = self.filter_var.get()
        self.ladda_lager_data(sok_term, filter_status)
    
    def visa_produktdetaljer(self, event):
        """Visa detaljer för vald produkt"""
        selection = self.tree.selection()
        if not selection:
            return
        
        item = self.tree.item(selection[0])
        sku = item['values'][0]
        
        # Hämta detaljer från databasen
        self.cursor.execute('''
            SELECT m.*, 
                   COALESCE(SUM(s.antal), 0) as totalt_antal,
                   COALESCE(SUM(s.reserverat), 0) as totalt_reserverat
            FROM masterdata m
            LEFT JOIN stock s ON m.sku = s.sku
            WHERE m.sku = ?
            GROUP BY m.sku
        ''', (sku,))
        
        produkt = self.cursor.fetchone()
        
        if produkt:
            # Skapa detaljfönster
            detail_window = tk.Toplevel(self.root)
            detail_window.title(f"Produktdetaljer: {sku}")
            detail_window.geometry("600x500")
            
            # Visa produktinfo
            info_frame = ttk.LabelFrame(detail_window, text="Produktinformation", padding=10)
            info_frame.pack(fill=tk.X, padx=10, pady=10)
            
            info_text = f"""
SKU: {produkt[0]}
Artikelnummer: {produkt[1]}
Benämning: {produkt[2]}
Kategori: {produkt[3]}
Underkategori: {produkt[4]}
Enhet: {produkt[5]}
Inköpspris: {produkt[6]:,.2f} SEK
Försäljningspris: {produkt[7]:,.2f} SEK
Leverantör: {produkt[9]}
Min nivå: {produkt[10]}
Max nivå: {produkt[11]}
Lead time: {produkt[12]} dagar
Totalt antal: {produkt[14]}
Reserverat: {produkt[15]}
Tillgängligt: {produkt[14] - produkt[15]}
            """
            
            ttk.Label(info_frame, text=info_text, justify=tk.LEFT).pack()
            
            # Lagerplatser
            lager_frame = ttk.LabelFrame(detail_window, text="Lagerplatser", padding=10)
            lager_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
            
            # Hämta lagerplatser
            self.cursor.execute('''
                SELECT lagerplats, hylla, rad, fack, batch, antal, senast_raknad
                FROM stock 
                WHERE sku = ? AND antal > 0
                ORDER BY lagerplats
            ''', (sku,))
            
            lagerplatser = self.cursor.fetchall()
            
            if lagerplatser:
                columns = ("Lagerplats", "Hylla", "Rad", "Fack", "Batch", "Antal", "Senast räknad")
                tree = ttk.Treeview(lager_frame, columns=columns, show="headings", height=5)
                
                for col in columns:
                    tree.heading(col, text=col)
                    tree.column(col, width=80)
                
                for row in lagerplatser:
                    tree.insert("", "end", values=row)
                
                scrollbar = ttk.Scrollbar(lager_frame, orient=tk.VERTICAL, command=tree.yview)
                tree.configure(yscrollcommand=scrollbar.set)
                
                tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
                scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
            else:
                ttk.Label(lager_frame, text="Inga lagerplatser registrerade").pack()
    
    def skapa_import_tab(self):
        """Skapa flik för Excel-import"""
        
        # Rubrik
        ttk.Label(self.tab_import, text="📁 IMPORTERA FRÅN EXCEL", 
                 font=("Arial", 16, "bold")).pack(pady=20)
        
        # Välj tabell att importera till
        ttk.Label(self.tab_import, text="Välj tabell att importera till:").pack(pady=10)
        
        self.import_tabell_var = tk.StringVar(value="masterdata")
        tabell_frame = ttk.Frame(self.tab_import)
        tabell_frame.pack(pady=10)
        
        tabeller = [
            ("Masterdata (produkter)", "masterdata"),
            ("Lager (stock)", "stock"),
            ("Inkommande ordrar", "inbound"),
            ("Utgående ordrar", "outbound")
        ]
        
        for i, (text, value) in enumerate(tabeller):
            rb = ttk.Radiobutton(tabell_frame, text=text, variable=self.import_tabell_var, 
                                value=value)
            rb.grid(row=i//2, column=i%2, padx=20, pady=5, sticky=tk.W)
        
        # Välj Excel-fil
        ttk.Label(self.tab_import, text="Välj Excel-fil:").pack(pady=15)
        
        file_frame = ttk.Frame(self.tab_import)
        file_frame.pack(pady=10)
        
        self.excel_file_var = tk.StringVar()
        ttk.Entry(file_frame, textvariable=self.excel_file_var, width=50).pack(side=tk.LEFT, padx=5)
        
        ttk.Button(file_frame, text="Bläddra...", 
                  command=lambda: self.bladdra_excel_fil()).pack(side=tk.LEFT, padx=5)
        
        # Välj Excel-ark (sheet)
        ttk.Label(self.tab_import, text="Välj ark i Excel-filen (lämna tomt för första arket):").pack(pady=10)
        self.sheet_name_var = tk.StringVar()
        ttk.Entry(self.tab_import, textvariable=self.sheet_name_var, width=30).pack(pady=5)
        
        # Förhandsgranska
        ttk.Button(self.tab_import, text="📋 Förhandsgranska data", 
                  command=self.forhandsgranska_excel).pack(pady=20)
        
        # Importera knapp
        ttk.Button(self.tab_import, text="🚀 IMPORTERA DATA", 
                  command=self.importera_excel, style="Accent.TButton").pack(pady=20)
        
        # Förhandsgranskningstext
        preview_frame = ttk.LabelFrame(self.tab_import, text="Förhandsgranskning", padding=10)
        preview_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)
        
        self.preview_text = tk.Text(preview_frame, height=10, wrap=tk.NONE)
        preview_scroll_y = ttk.Scrollbar(preview_frame, orient=tk.VERTICAL, command=self.preview_text.yview)
        preview_scroll_x = ttk.Scrollbar(preview_frame, orient=tk.HORIZONTAL, command=self.preview_text.xview)
        self.preview_text.configure(yscrollcommand=preview_scroll_y.set, xscrollcommand=preview_scroll_x.set)
        
        self.preview_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        preview_scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        preview_scroll_x.pack(side=tk.BOTTOM, fill=tk.X)
        
        # Mall-nedladdning
        ttk.Label(self.tab_import, text="💡 Tips: Ladda ner Excel-mallar:", 
                 font=("Arial", 10, "bold")).pack(pady=20)
        
        mall_frame = ttk.Frame(self.tab_import)
        mall_frame.pack(pady=10)
        
        for text, command in [
            ("📥 Masterdata mall", lambda: self.skapa_excel_mall("masterdata")),
            ("📥 Lager mall", lambda: self.skapa_excel_mall("stock")),
            ("📥 Inbound mall", lambda: self.skapa_excel_mall("inbound")),
            ("📥 Outbound mall", lambda: self.skapa_excel_mall("outbound"))
        ]:
            ttk.Button(mall_frame, text=text, command=command, width=15).pack(side=tk.LEFT, padx=5)
    
    def bladdra_excel_fil(self):
        """Öppna dialog för att välja Excel-fil"""
        filename = filedialog.askopenfilename(
            title="Välj Excel-fil",
            filetypes=[
                ("Excel files", "*.xlsx *.xls"),
                ("All files", "*.*")
            ]
        )
        if filename:
            self.excel_file_var.set(filename)
    
    def forhandsgranska_excel(self):
        """Förhandsgranska data från Excel-fil"""
        filename = self.excel_file_var.get()
        sheet_name = self.sheet_name_var.get() if self.sheet_name_var.get() else 0
        
        if not filename:
            messagebox.showwarning("Varning", "Välj en Excel-fil först!")
            return
        
        try:
            # Läs Excel-fil
            df = pd.read_excel(filename, sheet_name=sheet_name, nrows=20)
            
            # Rensa förhandsgranskning
            self.preview_text.delete(1.0, tk.END)
            
            # Visa information
            self.preview_text.insert(tk.END, f"Fil: {os.path.basename(filename)}\n")
            self.preview_text.insert(tk.END, f"Ark: {sheet_name if sheet_name else 'Första arket'}\n")
            self.preview_text.insert(tk.END, f"Antal rader: {len(df)} (visar max 20)\n")
            self.preview_text.insert(tk.END, f"Kolumner: {', '.join(df.columns)}\n")
            self.preview_text.insert(tk.END, "\n" + "="*50 + "\n\n")
            
            # Visa data som text
            self.preview_text.insert(tk.END, df.to_string())
            
        except Exception as e:
            messagebox.showerror("Fel", f"Kunde inte läsa Excel-filen:\n{str(e)}")
    
    def importera_excel(self):
        """Importera data från Excel-fil till databas"""
        filename = self.excel_file_var.get()
        tabell = self.import_tabell_var.get()
        sheet_name = self.sheet_name_var.get() if self.sheet_name_var.get() else 0
        
        if not filename:
            messagebox.showerror("Fel", "Välj en Excel-fil!")
            return
        
        if not os.path.exists(filename):
            messagebox.showerror("Fel", "Filen finns inte!")
            return
        
        try:
            # Läs hela Excel-filen
            df = pd.read_excel(filename, sheet_name=sheet_name)
            
            if df.empty:
                messagebox.showwarning("Varning", "Excel-filen är tom!")
                return
            
            # Ta bort eventuella tomma rader
            df = df.dropna(how='all')
            
            antal_importerade = 0
            
            if tabell == "masterdata":
                # Kräv minst SKU och benämning
                required_cols = ['sku', 'benamning']
                for col in required_cols:
                    if col not in df.columns:
                        messagebox.showerror("Fel", f"Excel-filen måste ha kolumnen '{col}'")
                        return
                
                # Importera till masterdata
                for _, rad in df.iterrows():
                    self.cursor.execute('''
                        INSERT OR REPLACE INTO masterdata 
                        (sku, artikelnummer, benamning, kategori, underkategori, enhet, 
                         inkopspris, forsaljningspris, leverantor, min_niva, max_niva)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        str(rad.get('sku', '')).strip(),
                        str(rad.get('artikelnummer', rad.get('sku', ''))).strip(),
                        str(rad.get('benamning', '')).strip(),
                        str(rad.get('kategori', '')).strip(),
                        str(rad.get('underkategori', '')).strip(),
                        str(rad.get('enhet', 'st')).strip(),
                        float(rad.get('inkopspris', 0)) if pd.notna(rad.get('inkopspris')) else 0,
                        float(rad.get('forsaljningspris', 0)) if pd.notna(rad.get('forsaljningspris')) else 0,
                        str(rad.get('leverantor', '')).strip(),
                        int(rad.get('min_niva', 5)) if pd.notna(rad.get('min_niva')) else 5,
                        int(rad.get('max_niva', 100)) if pd.notna(rad.get('max_niva')) else 100
                    ))
                    antal_importerade += 1
            
            elif tabell == "stock":
                # Kräv SKU och antal
                if 'sku' not in df.columns or 'antal' not in df.columns:
                    messagebox.showerror("Fel", "Excel-filen måste ha kolumnerna 'sku' och 'antal'")
                    return
                
                for _, rad in df.iterrows():
                    self.cursor.execute('''
                        INSERT OR REPLACE INTO stock 
                        (sku, lagerplats, hylla, rad, fack, batch, antal, senast_raknad)
                        VALUES (?, ?, ?, ?, ?, ?, ?, DATE('now'))
                    ''', (
                        str(rad.get('sku', '')).strip(),
                        str(rad.get('lagerplats', '')).strip(),
                        str(rad.get('hylla', '')).strip(),
                        str(rad.get('rad', '')).strip(),
                        str(rad.get('fack', '')).strip(),
                        str(rad.get('batch', '')).strip(),
                        int(rad.get('antal', 0)) if pd.notna(rad.get('antal')) else 0
                    ))
                    antal_importerade += 1
            
            elif tabell == "inbound":
                for _, rad in df.iterrows():
                    self.cursor.execute('''
                        INSERT INTO inbound 
                        (ordernummer, leverantorsorder, sku, leverantor, antal_orderat, 
                         enhetspris, forvantad_datum, status)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        str(rad.get('ordernummer', '')).strip(),
                        str(rad.get('leverantorsorder', '')).strip(),
                        str(rad.get('sku', '')).strip(),
                        str(rad.get('leverantor', '')).strip(),
                        int(rad.get('antal_orderat', 0)) if pd.notna(rad.get('antal_orderat')) else 0,
                        float(rad.get('enhetspris', 0)) if pd.notna(rad.get('enhetspris')) else 0,
                        rad.get('forvantad_datum') if pd.notna(rad.get('forvantad_datum')) else None,
                        str(rad.get('status', 'På väg')).strip()
                    ))
                    antal_importerade += 1
            
            elif tabell == "outbound":
                for _, rad in df.iterrows():
                    self.cursor.execute('''
                        INSERT INTO outbound 
                        (ordernummer, kundorder, sku, kund, kundnummer, antal_bestallt,
                         enhetspris, planerat_datum, ordrad_status)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        str(rad.get('ordernummer', '')).strip(),
                        str(rad.get('kundorder', '')).strip(),
                        str(rad.get('sku', '')).strip(),
                        str(rad.get('kund', '')).strip(),
                        str(rad.get('kundnummer', '')).strip(),
                        int(rad.get('antal_bestallt', 0)) if pd.notna(rad.get('antal_bestallt')) else 0,
                        float(rad.get('enhetspris', 0)) if pd.notna(rad.get('enhetspris')) else 0,
                        rad.get('planerat_datum') if pd.notna(rad.get('planerat_datum')) else None,
                        str(rad.get('ordrad_status', 'Öppen')).strip()
                    ))
                    antal_importerade += 1
            
            self.conn.commit()
            
            # Uppdatera lager-vyn
            if tabell in ["masterdata", "stock"]:
                self.ladda_lager_data()
            
            messagebox.showinfo("Import klar", f"Importerade {antal_importerade} rader till {tabell}")
            self.status_var.set(f"Importerade {antal_importerade} rader till {tabell}")
            
            # Rensa förhandsgranskning
            self.preview_text.delete(1.0, tk.END)
            
        except Exception as e:
            messagebox.showerror("Importfel", f"Kunde inte importera:\n{str(e)}")
    
    def skapa_excel_mall(self, mall_typ):
        """Skapa Excel-mall för import"""
        from openpyxl import Workbook
        from openpyxl.styles import Font
        
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"mall_{mall_typ}.xlsx"
        )
        
        if not filename:
            return
        
        wb = Workbook()
        ws = wb.active
        ws.title = mall_typ.capitalize()
        
        # Definiera kolumner för varje malltyp
        mallar = {
            "masterdata": [
                "sku", "artikelnummer", "benamning", "kategori", "underkategori",
                "enhet", "inkopspris", "forsaljningspris", "leverantor",
                "min_niva", "max_niva", "lead_time"
            ],
            "stock": [
                "sku", "lagerplats", "hylla", "rad", "fack", "batch",
                "serienummer", "antal", "utgangsdatum", "kvalitetskod"
            ],
            "inbound": [
                "ordernummer", "leverantorsorder", "sku", "leverantor",
                "antal_orderat", "enhetspris", "forvantad_datum", "status"
            ],
            "outbound": [
                "ordernummer", "kundorder", "sku", "kund", "kundnummer",
                "antal_bestallt", "enhetspris", "planerat_datum", "ordrad_status"
            ]
        }
        
        if mall_typ not in mallar:
            messagebox.showerror("Fel", f"Okänd malltyp: {mall_typ}")
            return
        
        # Lägg till kolumnnamn
        for col_num, kolumn in enumerate(mallar[mall_typ], 1):
            cell = ws.cell(row=1, column=col_num, value=kolumn)
            cell.font = Font(bold=True)
        
        # Lägg till exempeldata
        if mall_typ == "masterdata":
            exempel_data = [
                ["SKU001", "ART001", "Dator Laptop", "Elektronik", "Datorer", 
                 "st", 10000.00, 15000.00, "Dell", 5, 50, 7],
                ["SKU002", "ART002", "Mus USB", "Elektronik", "Tillbehör", 
                 "st", 200.00, 299.00, "Logitech", 20, 200, 5]
            ]
        elif mall_typ == "stock":
            exempel_data = [
                ["SKU001", "LAGER-A", "A", "01", "01", "BATCH2024-01", 
                 "SN001", 25, "2025-12-31", "OK"],
                ["SKU002", "LAGER-B", "B", "03", "02", "BATCH2024-02", 
                 "SN002", 150, "2025-06-30", "OK"]
            ]
        elif mall_typ == "inbound":
            exempel_data = [
                ["PO-2024-001", "LO-2024-001", "SKU001", "Dell", 
                 10, 10000.00, "2024-02-15", "På väg"],
                ["PO-2024-002", "LO-2024-002", "SKU002", "Logitech", 
                 50, 200.00, "2024-02-20", "På väg"]
            ]
        elif mall_typ == "outbound":
            exempel_data = [
                ["SO-2024-001", "KO-2024-001", "SKU001", "ABC Bolag", "KUND001", 
                 5, 15000.00, "2024-02-10", "Öppen"],
                ["SO-2024-002", "KO-2024-002", "SKU002", "XYZ AB", "KUND002", 
                 20, 299.00, "2024-02-12", "Öppen"]
            ]
        
        for row_num, rad in enumerate(exempel_data, 2):
            for col_num, värde in enumerate(rad, 1):
                ws.cell(row=row_num, column=col_num, value=värde)
        
        # Justera kolumnbredder
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Spara fil
        wb.save(filename)
        messagebox.showinfo("Mall skapad", f"Excel-mall sparad som:\n{filename}")
    
    def skapa_rapporter_tab(self):
        """Skapa flik för rapporter"""
        
        ttk.Label(self.tab_rapporter, text="📊 RAPPORTER OCH EXPORT", 
                 font=("Arial", 16, "bold")).pack(pady=20)
        
        # Rapportknappar
        rapport_frame = ttk.Frame(self.tab_rapporter)
        rapport_frame.pack(pady=20)
        
        rapporter = [
            ("📈 Lagerstatus", self.export_lagerstatus),
            ("⚠️ Overstock", self.export_overstock),
            ("📉 Lågt lager", self.export_lagt_lager),
            ("💀 Deadstock", self.export_deadstock),
            ("📦 Allt lager", self.export_allt_lager),
            ("📋 Återbeställningslista", self.export_aterbestallning),
            ("💰 Lagervärde per kategori", self.export_lagervarde),
            ("📅 Veckorapport", self.export_veckorapport)
        ]
        
        for i, (text, command) in enumerate(rapporter):
            row, col = divmod(i, 2)
            ttk.Button(rapport_frame, text=text, command=command, width=25).grid(
                row=row, column=col, padx=10, pady=10
            )
        
        # Anpassad rapport
        ttk.Label(self.tab_rapporter, text="Anpassad rapport:", 
                 font=("Arial", 12, "bold")).pack(pady=(30, 10))
        
        custom_frame = ttk.Frame(self.tab_rapporter)
        custom_frame.pack(pady=10)
        
        ttk.Button(custom_frame, text="📊 Skapa anpassad rapport", 
                  command=self.skapa_anpassad_rapport).pack(side=tk.LEFT, padx=5)
        
        # Export av all data
        ttk.Label(self.tab_rapporter, text="Fullständig export:", 
                 font=("Arial", 12, "bold")).pack(pady=(30, 10))
        
        ttk.Button(self.tab_rapporter, text="💾 Exportera ALL data till Excel (alla tabeller)", 
                  command=self.export_all_data, style="Accent.TButton").pack(pady=10)
    
    def export_lagerstatus(self):
        """Exportera lagerstatus till Excel"""
        self.export_rapport_till_excel('''
            SELECT * FROM v_lagerstatus 
            ORDER BY sku
        ''', "lagerstatus")
    
    def export_overstock(self):
        """Exportera overstock till Excel"""
        self.export_rapport_till_excel('''
            SELECT sku, artikelnummer, benamning, kategori, totalt_antal, 
                   max_niva, fyllnadsgrad, lager_varde
            FROM v_lagerstatus 
            WHERE status LIKE '%OVERSTOCK%'
            ORDER BY fyllnadsgrad DESC
        ''', "overstock")
    
    def export_lagt_lager(self):
        """Exportera lågt lager till Excel"""
        self.export_rapport_till_excel('''
            SELECT sku, artikelnummer, benamning, kategori, tillgangligt, 
                   min_niva, (min_niva - tillgangligt) as brist, lager_varde
            FROM v_lagerstatus 
            WHERE status LIKE '%LÅGT%' OR tillgangligt < min_niva
            ORDER BY brist DESC
        ''', "lagt_lager")
    
    def export_deadstock(self):
        """Exportera deadstock till Excel"""
        self.export_rapport_till_excel('''
            SELECT sku, artikelnummer, benamning, kategori, tillgangligt, 
                   dagar_sen_sista_rorelse, lager_varde
            FROM v_lagerstatus 
            WHERE status LIKE '%DEADSTOCK%' OR dagar_sen_sista_rorelse > 90
            ORDER BY dagar_sen_sista_rorelse DESC
        ''', "deadstock")
    
    def export_allt_lager(self):
        """Exportera allt lager till Excel"""
        self.export_rapport_till_excel('''
            SELECT s.sku, m.benamning, m.kategori, s.lagerplats, s.hylla, 
                   s.rad, s.fack, s.batch, s.antal, s.reserverat, s.tillgangligt,
                   s.senast_raknad
            FROM stock s
            JOIN masterdata m ON s.sku = m.sku
            WHERE s.antal > 0
            ORDER BY s.lagerplats, s.hylla, s.rad, s.fack
        ''', "allt_lager_detaljerat")
    
    def export_aterbestallning(self):
        """Exportera återbeställningslista till Excel"""
        self.export_rapport_till_excel('''
            SELECT sku, artikelnummer, benamning, kategori, tillgangligt, 
                   min_niva, (min_niva - tillgangligt) as att_bestalla,
                   leverantor, lead_time
            FROM v_lagerstatus vs
            JOIN masterdata m ON vs.sku = m.sku
            WHERE tillgangligt < min_niva
            ORDER BY (min_niva - tillgangligt) DESC
        ''', "aterbestallningslista")
    
    def export_lagervarde(self):
        """Exportera lager värde per kategori till Excel"""
        self.export_rapport_till_excel('''
            SELECT kategori, 
                   COUNT(DISTINCT sku) as antal_produkter,
                   SUM(totalt_antal) as totalt_antal,
                   SUM(lager_varde) as totalt_varde,
                   ROUND(AVG(fyllnadsgrad), 1) as genomsnittlig_fyllnadsgrad
            FROM v_lagerstatus 
            GROUP BY kategori
            ORDER BY totalt_varde DESC
        ''', "lagervarde_per_kategori")
    
    def export_veckorapport(self):
        """Exportera veckorapport till Excel"""
        datum = datetime.now().strftime("%Y-%m-%d")
        self.export_rapport_till_excel(f'''
            SELECT 
                'Lagerstatus {datum}' as rapport_datum,
                COUNT(DISTINCT sku) as totalt_antal_sku,
                SUM(totalt_antal) as totalt_lagerantal,
                SUM(lager_varde) as totalt_lagervarde,
                SUM(CASE WHEN status LIKE '%LÅGT%' THEN 1 ELSE 0 END) as antal_lagt_lager,
                SUM(CASE WHEN status LIKE '%OVERSTOCK%' THEN 1 ELSE 0 END) as antal_overstock,
                SUM(CASE WHEN status LIKE '%DEADSTOCK%' THEN 1 ELSE 0 END) as antal_deadstock,
                SUM(CASE WHEN status LIKE '%SLUT%' THEN 1 ELSE 0 END) as antal_slut
            FROM v_lagerstatus
        ''', f"veckorapport_{datum}")
    
    def export_rapport_till_excel(self, query, rapport_namn):
        """Generisk funktion för att exportera rapport till Excel"""
        try:
            # Kör SQL-frågan
            df = pd.read_sql_query(query, self.conn)
            
            if df.empty:
                messagebox.showinfo("Ingen data", f"Ingen data att exportera för {rapport_namn}")
                return
            
            # Välj filnamn
            datum = datetime.now().strftime("%Y%m%d_%H%M")
            filename = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile=f"{rapport_namn}_{datum}.xlsx"
            )
            
            if not filename:
                return
            
            # Exportera till Excel
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name=rapport_namn[:30], index=False)
                
                # Lägg till formatering
                workbook = writer.book
                worksheet = writer.sheets[rapport_namn[:30]]
                
                # Gör rubriker fetstilta
                for cell in worksheet[1]:
                    cell.font = openpyxl.styles.Font(bold=True)
                
                # Justera kolumnbredder
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length:
                                max_length = cell_length
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
            
            messagebox.showinfo("Export klar", f"Rapport exporterad till:\n{filename}")
            self.status_var.set(f"Exporterade {rapport_namn} till Excel")
            
        except Exception as e:
            messagebox.showerror("Exportfel", f"Kunde inte exportera:\n{str(e)}")
    
    def export_all_data(self):
        """Exportera all data från alla tabeller till Excel"""
        try:
            datum = datetime.now().strftime("%Y%m%d_%H%M")
            filename = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile=f"lager_full_export_{datum}.xlsx"
            )
            
            if not filename:
                return
            
            # Skapa Excel-fil med flikar
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                # Exportera varje tabell
                tabeller = ['masterdata', 'stock', 'inbound', 'outbound', 'v_lagerstatus']
                
                for tabell in tabeller:
                    try:
                        df = pd.read_sql_query(f"SELECT * FROM {tabell}", self.conn)
                        if not df.empty:
                            sheet_name = tabell[:30]  # Excel max 31 tecken för fliknamn
                            df.to_excel(writer, sheet_name=sheet_name, index=False)
                    except:
                        pass  # Skip tabeller som inte finns
                
                # Lägg till en sammanfattningsflik
                summary_data = {
                    'Tabell': ['masterdata', 'stock', 'inbound', 'outbound', 'v_lagerstatus'],
                    'Beskrivning': ['Produktinformation', 'Aktuellt lager', 'Inkommande ordrar', 
                                   'Utgående ordrar', 'Lagerstatus vy'],
                    'Antal rader': []
                }
                
                for tabell in summary_data['Tabell']:
                    try:
                        self.cursor.execute(f"SELECT COUNT(*) FROM {tabell}")
                        count = self.cursor.fetchone()[0]
                        summary_data['Antal rader'].append(count)
                    except:
                        summary_data['Antal rader'].append(0)
                
                df_summary = pd.DataFrame(summary_data)
                df_summary.to_excel(writer, sheet_name='Sammanfattning', index=False)
            
            messagebox.showinfo("Export klar", f"All data exporterad till:\n{filename}")
            self.status_var.set("All data exporterad till Excel")
            
        except Exception as e:
            messagebox.showerror("Exportfel", f"Kunde inte exportera all data:\n{str(e)}")
    
    def skapa_anpassad_rapport(self):
        """Dialog för att skapa anpassad rapport"""
        dialog = tk.Toplevel(self.root)
        dialog.title("Anpassad rapport")
        dialog.geometry("600x400")
        
        ttk.Label(dialog, text="Skapa anpassad SQL-fråga:", 
                 font=("Arial", 12, "bold")).pack(pady=20)
        
        # SQL-frågeruta
        sql_text = tk.Text(dialog, height=10, width=70)
        sql_text.pack(pady=10, padx=20)
        sql_text.insert(tk.END, "SELECT sku, benamning, kategori, totalt_antal FROM v_lagerstatus WHERE 1=1")
        
        # Kör och exportknappar
        btn_frame = ttk.Frame(dialog)
        btn_frame.pack(pady=20)
        
        def kor_fraga():
            query = sql_text.get("1.0", tk.END).strip()
            if not query:
                messagebox.showwarning("Varning", "Ange en SQL-fråga!")
                return
            
            try:
                # Testkör frågan
                df = pd.read_sql_query(query, self.conn)
                
                # Visa förhandsgranskning
                preview_window = tk.Toplevel(dialog)
                preview_window.title("Förhandsgranskning")
                preview_window.geometry("800x400")
                
                ttk.Label(preview_window, text=f"Resultat: {len(df)} rader", 
                         font=("Arial", 12)).pack(pady=10)
                
                # Visa data i trädvy
                columns = list(df.columns)
                tree = ttk.Treeview(preview_window, columns=columns, show="headings", height=15)
                
                for col in columns:
                    tree.heading(col, text=col)
                    tree.column(col, width=100)
                
                for _, row in df.iterrows():
                    tree.insert("", "end", values=list(row))
                
                scrollbar = ttk.Scrollbar(preview_window, orient=tk.VERTICAL, command=tree.yview)
                tree.configure(yscrollcommand=scrollbar.set)
                
                tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=10, pady=10)
                scrollbar.pack(side=tk.RIGHT, fill=tk.Y, pady=10)
                
                def export_resultat():
                    self.export_anpassad_rapport(query)
                
                ttk.Button(preview_window, text="📊 Exportera till Excel", 
                          command=export_resultat).pack(pady=10)
                
            except Exception as e:
                messagebox.showerror("SQL-fel", f"Ogiltig SQL-fråga:\n{str(e)}")
        
        def export_fraga():
            query = sql_text.get("1.0", tk.END).strip()
            if not query:
                messagebox.showwarning("Varning", "Ange en SQL-fråga!")
                return
            self.export_anpassad_rapport(query)
        
        ttk.Button(btn_frame, text="▶️ Kör fråga", command=kor_fraga).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="📊 Exportera direkt", command=export_fraga).pack(side=tk.LEFT, padx=5)
    
    def export_anpassad_rapport(self, query):
        """Exportera anpassad rapport till Excel"""
        try:
            df = pd.read_sql_query(query, self.conn)
            
            if df.empty:
                messagebox.showinfo("Ingen data", "Ingen data att exportera")
                return
            
            datum = datetime.now().strftime("%Y%m%d_%H%M")
            filename = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile=f"anpassad_rapport_{datum}.xlsx"
            )
            
            if filename:
                df.to_excel(filename, index=False)
                messagebox.showinfo("Export klar", f"Rapport exporterad till:\n{filename}")
                self.status_var.set("Anpassad rapport exporterad")
                
        except Exception as e:
            messagebox.showerror("Exportfel", f"Kunde inte exportera:\n{str(e)}")
    
    def skapa_admin_tab(self):
        """Skapa administrationsflik"""
        
        ttk.Label(self.tab_admin, text="⚙️ ADMINISTRATION", 
                 font=("Arial", 16, "bold")).pack(pady=20)
        
        # Databashantering
        db_frame = ttk.LabelFrame(self.tab_admin, text="Databashantering", padding=15)
        db_frame.pack(fill=tk.X, padx=20, pady=10)
        
        ttk.Button(db_frame, text="🔄 Optimera databas", 
                  command=self.optimera_databas).pack(pady=5)
        ttk.Button(db_frame, text="💾 Skapa backup", 
                  command=self.skapa_backup).pack(pady=5)
        ttk.Button(db_frame, text="🗑️ Rensa gamla transaktioner", 
                  command=self.rensa_gamla_transaktioner).pack(pady=5)
        ttk.Button(db_frame, text="🔍 Databasstatistik", 
                  command=self.visa_databasstatistik).pack(pady=5)
        
        # Systeminställningar
        settings_frame = ttk.LabelFrame(self.tab_admin, text="Systeminställningar", padding=15)
        settings_frame.pack(fill=tk.X, padx=20, pady=10)
        
        ttk.Button(settings_frame, text="⚙️ Inställningar", 
                  command=self.visa_installningar).pack(pady=5)
        ttk.Button(settings_frame, text="🆘 Hjälp", 
                  command=self.visa_hjalp).pack(pady=5)
        ttk.Button(settings_frame, text="ℹ️ Om programmet", 
                  command=self.visa_om).pack(pady=5)
    
    def optimera_databas(self):
        """Optimera SQLite-databasen"""
        try:
            self.cursor.execute("VACUUM")
            self.conn.commit()
            messagebox.showinfo("Optimering", "Databas optimerad!")
            self.status_var.set("Databas optimerad")
        except Exception as e:
            messagebox.showerror("Fel", f"Kunde inte optimera databas:\n{str(e)}")
    
    def skapa_backup(self):
        """Skapa backup av databasen"""
        try:
            import shutil
            from datetime import datetime
            
            datum = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_file = f"lager_backup_{datum}.db"
            
            # Kopiera databasfilen
            shutil.copy2("lager_excel.db", backup_file)
            
            messagebox.showinfo("Backup", f"Backup skapad:\n{backup_file}")
            self.status_var.set(f"Backup skapad: {backup_file}")
            
        except Exception as e:
            messagebox.showerror("Fel", f"Kunde inte skapa backup:\n{str(e)}")
    
    def rensa_gamla_transaktioner(self):
        """Rensa gamla transaktioner"""
        try:
            self.cursor.execute("DELETE FROM transaktioner WHERE date(tidpunkt) < date('now', '-90 days')")
            antal_rader = self.cursor.rowcount
            self.conn.commit()
            
            messagebox.showinfo("Rensning", f"Raderade {antal_rader} gamla transaktioner")
            self.status_var.set(f"Raderade {antal_rader} gamla transaktioner")
            
        except Exception as e:
            messagebox.showerror("Fel", f"Kunde inte rensa:\n{str(e)}")
    
    def visa_databasstatistik(self):
        """Visa databasstatistik"""
        try:
            statistik = {}
            
            # Räkna poster i varje tabell
            for tabell in ['masterdata', 'stock', 'inbound', 'outbound']:
                self.cursor.execute(f"SELECT COUNT(*) FROM {tabell}")
                statistik[tabell] = self.cursor.fetchone()[0]
            
            # Total lagerantal
            self.cursor.execute("SELECT SUM(antal) FROM stock")
            totalt_antal = self.cursor.fetchone()[0] or 0
            
            # Totalt lager värde
            self.cursor.execute('''
                SELECT COALESCE(SUM(s.antal * m.forsaljningspris), 0)
                FROM stock s
                JOIN masterdata m ON s.sku = m.sku
            ''')
            totalt_varde = self.cursor.fetchone()[0] or 0
            
            # Visa statistik
            statistik_text = f"""
Databasstatistik:
----------------
Masterdata: {statistik.get('masterdata', 0)} produkter
Lagerposter: {statistik.get('stock', 0)} stycken
Inkommande ordrar: {statistik.get('inbound', 0)} stycken
Utgående ordrar: {statistik.get('outbound', 0)} stycken
Totalt lagerantal: {totalt_antal} stycken
Totalt lagervärde: {totalt_varde:,.0f} SEK
Databasfil: lager_excel.db
            """
            
            messagebox.showinfo("Databasstatistik", statistik_text)
            
        except Exception as e:
            messagebox.showerror("Fel", f"Kunde inte hämta statistik:\n{str(e)}")
    
    def visa_installningar(self):
        """Visa inställningar"""
        messagebox.showinfo("Inställningar", "Inställningar kommer i nästa version!")
    
    def visa_hjalp(self):
        """Visa hjälp"""
        hjalp_text = """
LAGERHANTERINGSSYSTEM - HJÄLP
-----------------------------

1. IMPORT AV DATA:
   - Använd Excel-mallarna för att förbereda data
   - Välj rätt tabell (masterdata, stock, inbound, outbound)
   - Välj Excel-fil och klicka "Importera"

2. LAGERÖVERSIKT:
   - Sök på SKU, artikelnummer eller benämning
   - Filtrera på status (Lågt, Overstock, Deadstock, Slut)
   - Dubbelklick på en produkt för detaljer

3. RAPPORTER:
   - Klicka på rapportknapparna för att exportera till Excel
   - Använd "Anpassad rapport" för egna SQL-frågor

4. ADMINISTRATION:
   - Skapa regelbundna backups
   - Optimera databasen periodvis

SUPPORT:
   Om du har frågor, kontakta IT-support.
        """
        messagebox.showinfo("Hjälp", hjalp_text)
    
    def visa_om(self):
        """Visa information om programmet"""
        om_text = """
Lagerhanteringssystem - Excel Version
Version: 2.0
Byggd med: Python, Tkinter, SQLite, Pandas
Datum: 2024

Funktioner:
- Excel-import/export
- Lagerstatus med färgkodning
- Overstock/Deadstock rapporter
- Fullständig rapportering
- Databashantering

Utvecklat för: Effektiv lagerstyrning
        """
        messagebox.showinfo("Om programmet", om_text)

# Kör appen
if __name__ == "__main__":
    # Styling
    style = ttk.Style()
    style.theme_use('clam')
    
    # Definiera en accentfärg för viktiga knappar
    style.configure("Accent.TButton", font=("Arial", 10, "bold"), padding=10)
    
    print("=" * 60)
    print("LAGERHANTERINGSSYSTEM - EXCEL VERSION")
    print("=" * 60)
    print("Databas: lager_excel.db")
    print("Startar...")
    
    app = LagerSystemExcel()